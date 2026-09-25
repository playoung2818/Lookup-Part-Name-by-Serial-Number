"""Create only the borrowing table and optionally seed an empty table from Excel.

Usage: python setup_borrow_log.py [path-to-workbook.xlsx]
The workbook is read only. Existing borrowing records are never overwritten.
"""
import argparse
from datetime import datetime

from sqlalchemy import select

from app import create_app
from models.models import db, WarehouseBorrowLog
from routes.Borrow_Log_Routes import FIELDS, parse_record


def workbook_records(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    records = []
    try:
        sheet = workbook['Borrowing Log']
        # Workbook columns are independent of the web form's display order.
        columns = {'borrower': 2, 'item_number': 3, 'serial_number': 4, 'borrowed_qty': 5,
                   'warehouse': 6, 'purpose': 7, 'borrow_date': 8, 'returned_qty': 9,
                   'return_date': 11, 'issued_by': 12}
        expected = ['Borrow ID', 'Borrower', 'Item #', 'Serial Number', 'Borrowed Qty', 'Warehouse',
                    'Purpose', 'Borrow Date', 'Returned Qty', 'Borrowed Duration (Days)', 'Return Date', 'Issued By', 'Status']
        if [sheet.cell(7, n).value for n in range(1, 14)] != expected:
            raise ValueError('Workbook columns do not match the borrowing log template.')
        for row in sheet.iter_rows(min_row=8):
            if not row[1].value:
                continue
            form = {}
            for name, column in columns.items():
                value = row[column - 1].value
                if isinstance(value, datetime):
                    value = value.date().isoformat()
                if value is None:
                    value = 0 if name == 'returned_qty' else ''
                form[name] = str(value)
            values, errors = parse_record(form)
            if errors:
                raise ValueError(f'Row {row[1].row}: {"; ".join(errors)}')
            # Preserve the workbook row suffix, including gaps between records.
            records.append(WarehouseBorrowLog(id=row[1].row - 7, **values))
    finally:
        workbook.close()
    return records


def setup(path=None):
    records = workbook_records(path) if path else []
    WarehouseBorrowLog.__table__.create(db.engine, checkfirst=True)
    if not records:
        return 'warehouse_borrow_log table is ready; no records imported.'
    # Serialize setup runs; never seed a table that already contains user records.
    if db.engine.dialect.name == 'postgresql':
        db.session.execute(db.text('LOCK TABLE warehouse_borrow_log IN EXCLUSIVE MODE'))
    if db.session.execute(select(WarehouseBorrowLog.id).limit(1)).first():
        db.session.rollback()
        return 'warehouse_borrow_log already contains records; import skipped.'
    db.session.add_all(records)
    db.session.flush()
    if db.engine.dialect.name == 'postgresql':
        db.session.execute(db.text("SELECT setval(pg_get_serial_sequence('warehouse_borrow_log', 'id'), (SELECT max(id) FROM warehouse_borrow_log))"))
    db.session.commit()
    return f'Imported {len(records)} records into warehouse_borrow_log.'


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', nargs='?')
    args = parser.parse_args()
    with create_app().app_context():
        print(setup(args.workbook))
