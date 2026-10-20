import io
import re
import unittest
from datetime import date
from unittest.mock import MagicMock, patch

from openpyxl import Workbook
from sqlalchemy import MetaData, Table, Column, String, Date, Float, Integer
from sqlalchemy.exc import SQLAlchemyError

from app import create_app
from functions.receiving_import import read_upload, classify, import_records, KEYS, build_review, apply_review, snapshot


def workbook(rows, headers=None, sheet='Receiving'):
    book = Workbook()
    book.active.title = sheet
    book.active.append(headers or ['Date', 'Inv# ', 'Box #', 'POD#', 'Part#', 'SN#', 'QTY'])
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def excel_row(serial='SN-1', qty=1):
    return [date(2026, 10, 8), 'INV-1', '', 'POD-1', 'PART-1', serial, qty]


def records(*rows):
    return read_upload(workbook(rows))[0]


def connection_with(existing):
    connection = MagicMock()
    connection.execute.return_value.mappings.return_value = existing
    return connection


def table():
    # In-memory SQL metadata only; these tests never create a database table.
    return Table('receiving_log', MetaData(), Column('id', Integer), Column('Reference', String), *(Column(key, Date if key == 'entry_date' else Float if key == 'quantity' else String) for key in KEYS))


class ReceivingValidationTests(unittest.TestCase):
    def test_read_trim_placeholders_default_quantity(self):
        result, errors, skipped = read_upload(workbook([excel_row(' SN-1 ', None), excel_row('NA'), excel_row('nan'), excel_row('')]))
        self.assertEqual(errors, [])
        self.assertEqual(skipped, 3)
        self.assertEqual(result[0]['serial_number'], 'SN-1')
        self.assertEqual(result[0]['quantity'], 1)
        self.assertEqual(result[0]['entry_date'], '2026-10-08')

    def test_bad_values_block_not_coerce(self):
        for quantity in ['bad', 0, -1, 'nan', 'inf', '=1+2']:
            with self.subTest(quantity=quantity):
                self.assertTrue(read_upload(workbook([excel_row(qty=quantity)]))[1])
        row = excel_row()
        row[0] = 'not a date'
        row[4] = ''
        self.assertEqual(len(read_upload(workbook([row]))[1]), 1)

    def test_first_worksheet_can_have_any_name(self):
        result, errors, skipped = read_upload(workbook([excel_row()], sheet='Sheet1'))
        self.assertEqual(len(result), 1)
        self.assertEqual(errors, [])

    def test_first_worksheet_is_used_not_named_receiving_or_active_sheet(self):
        book = Workbook()
        book.active.title = 'First'
        book.active.append(['Date', 'Inv#', 'Box #', 'POD#', 'Part#', 'SN#', 'QTY'])
        book.active.append(excel_row('FIRST-SN'))
        second = book.create_sheet('Receiving')
        second.append(['Wrong headers'])
        book.active = 1
        output = io.BytesIO()
        book.save(output)
        book.close()
        result, errors, skipped = read_upload(output.getvalue())
        self.assertEqual(result[0]['serial_number'], 'FIRST-SN')
        self.assertEqual(errors, [])

    def test_headers_are_case_insensitive_and_extra_columns_allowed(self):
        headers = ['Date', 'INV#', 'Box #', 'POD#', 'Part#', 'SN#', 'Qty', 'Reference', 'Column1']
        result, errors, skipped = read_upload(workbook([excel_row()], headers=headers))
        self.assertEqual(errors, [])
        self.assertEqual(result[0]['invoice_number'], 'INV-1')
        self.assertEqual(result[0]['quantity'], 1)

    def test_duplicate_header_with_different_case_is_rejected(self):
        headers = ['Date', 'Inv#', 'Box #', 'POD#', 'Part#', 'SN#', 'QTY', 'INV#']
        with self.assertRaisesRegex(ValueError, "exactly one 'Inv#'"):
            read_upload(workbook([], headers=headers))

    def test_missing_header_still_rejected(self):
        with self.assertRaises(ValueError):
            read_upload(workbook([], headers=['SN#']))

    def test_exact_upload_duplicates_and_case_normalization(self):
        rows = records(excel_row(), excel_row(' sn-1 '))
        pending, errors, skipped = classify(connection_with([]), table(), rows)
        self.assertEqual((len(pending), errors, skipped), (1, [], 1))

    def test_upload_conflicting_duplicate_blocks(self):
        rows = records(excel_row(), excel_row('sn-1', 2))
        self.assertTrue(classify(connection_with([]), table(), rows)[1])

    def test_existing_exact_skipped_but_conflicts_block(self):
        rows = records(excel_row())
        existing = dict(rows[0], entry_date=date(2026, 10, 8))
        self.assertEqual(classify(connection_with([existing]), table(), rows), ([], [], 1))
        existing['quantity'] = 3
        self.assertTrue(classify(connection_with([existing]), table(), rows)[1])
        self.assertTrue(classify(connection_with([existing, existing]), table(), rows)[1])

    def test_commit_rechecks_under_postgres_write_lock(self):
        engine = MagicMock()
        connection = engine.connect.return_value.__enter__.return_value
        connection.dialect.name = 'postgresql'
        with patch('functions.receiving_import.receiving_table', return_value=table()), patch('functions.receiving_import.classify', return_value=(records(excel_row()), [], 0)):
            self.assertEqual(import_records(engine, records(excel_row())), (1, 0))
        statements = [str(call.args[0]) for call in connection.execute.call_args_list]
        self.assertIn('SHARE ROW EXCLUSIVE', statements[1])
        connection.commit.assert_called_once()

    def test_conflict_or_insert_error_rolls_back(self):
        for errors in [['serial conflict'], []]:
            engine = MagicMock()
            connection = engine.connect.return_value.__enter__.return_value
            connection.dialect.name = 'sqlite'
            if not errors:
                connection.execute.side_effect = SQLAlchemyError('insert failure')
            with patch('functions.receiving_import.receiving_table', return_value=table()), patch('functions.receiving_import.classify', return_value=(records(excel_row()), errors, 0)):
                with self.assertRaises((ValueError, SQLAlchemyError)):
                    import_records(engine, records(excel_row()))
            connection.rollback.assert_called_once()
            connection.commit.assert_not_called()


class ReceivingRouteTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app({'TESTING': True, 'SECRET_KEY': 'test', 'SQLALCHEMY_DATABASE_URI': 'sqlite://'})
        self.client = self.app.test_client()
        self.assertEqual(self.client.get('/receiving').status_code, 200)
        with self.client.session_transaction() as session:
            self.csrf = session['receiving_csrf']

    def tearDown(self):
        from models.models import db
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()

    def preview(self, rows):
        review = [{'key': 'SN-1', 'uploaded': records(*rows), 'existing': [], 'snapshot': snapshot([]), 'kind': 'new'}]
        with patch('routes.Receiving_Routes.receiving_table', return_value=table()), patch('routes.Receiving_Routes.build_review', return_value=(review, 0)):
            return self.client.post('/receiving', data={'csrf_token': self.csrf, 'file': (io.BytesIO(workbook(rows)), 'receiving.xlsm')})

    def test_preview_never_writes_and_confirmation_required(self):
        with patch('routes.Receiving_Routes.apply_review', return_value={'inserted': 1, 'updated': 0, 'deleted': 0, 'untouched': 0}) as save:
            response = self.preview([excel_row()])
            self.assertEqual(response.status_code, 200)
            save.assert_not_called()
            token = re.search(rb'name="preview_token" value="([^"]+)"', response.data).group(1).decode()
            response = self.client.post('/receiving', data={'csrf_token': self.csrf, 'action': 'import', 'preview_token': token})
            self.assertIn(b'1 rows inserted', response.data)
            save.assert_called_once()

    def test_csrf_and_tampered_preview_rejected(self):
        self.assertEqual(self.client.post('/receiving', data={}).status_code, 400)
        with patch('routes.Receiving_Routes.apply_review') as save:
            response = self.client.post('/receiving', data={'csrf_token': self.csrf, 'action': 'import', 'preview_token': 'modified'})
            self.assertEqual(response.status_code, 400)
            save.assert_not_called()

    def test_all_conflicts_render_with_choices_and_escaped_notes(self):
        review = []
        for i in range(142):
            uploaded = records(excel_row(f'SN-{i}'))
            existing = [dict(uploaded[0], id=i * 2, entry_date=date(2026, 10, 8), Reference='<script>bad()</script>'),
                        dict(uploaded[0], id=i * 2 + 1, Reference='note')]
            review.append({'key': f'SN-{i}', 'uploaded': uploaded, 'existing': existing,
                           'snapshot': snapshot(existing), 'kind': 'conflict'})
        with patch('routes.Receiving_Routes.receiving_table', return_value=table()), patch('routes.Receiving_Routes.build_review', return_value=(review, 18029)):
            response = self.client.post('/receiving', data={'csrf_token': self.csrf, 'file': (io.BytesIO(workbook([excel_row()])), 'receiving.xlsx')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data.count(b'class="serial-group"'), 142)
        self.assertIn(b'name="choice_141"', response.data)
        self.assertIn(b'&lt;script&gt;bad()&lt;/script&gt;', response.data)
        self.assertNotIn(b'<script>bad()', response.data)
        self.assertIn(b'name="preview_token"', response.data)
        self.assertIn(b'name="allow_delete"', response.data)

    def test_validation_errors_have_no_confirm_button(self):
        response = self.preview([excel_row(qty=-1)])
        self.assertIn(b'Import blocked', response.data)
        self.assertNotIn(b'name="preview_token"', response.data)


class ReceivingReviewTests(unittest.TestCase):
    def setUp(self):
        self.uploaded = records(excel_row())
        self.db_rows = [dict(self.uploaded[0], id=10, Reference='keep note'),
                        dict(self.uploaded[0], id=20, Reference='other note')]
        for row in self.db_rows:
            row.pop('_row', None)
        self.group = {'key': 'SN-1', 'uploaded': self.uploaded, 'existing': self.db_rows,
                      'snapshot': snapshot(self.db_rows), 'kind': 'conflict'}
        self.engine = MagicMock()
        self.connection = self.engine.connect.return_value.__enter__.return_value
        self.connection.dialect.name = 'postgresql'
        self.connection.execute.return_value.rowcount = 1

    def apply(self, choice, allow=False, current=None):
        with patch('functions.receiving_import.receiving_table', return_value=table()), patch('functions.receiving_import.database_groups', return_value={'SN-1': self.db_rows if current is None else current}):
            return apply_review(self.engine, [self.group], {'0': choice}, allow)

    def test_existing_duplicates_become_review_not_errors(self):
        review, skipped = build_review(connection_with(self.db_rows), table(), self.uploaded * 2)
        self.assertEqual(len(review), 1)
        self.assertEqual(len(review[0]['existing']), 2)
        self.assertEqual(review[0]['kind'], 'conflict')
        self.assertEqual(skipped, 1)

    def test_uploaded_conflicting_variants_are_all_available(self):
        review, skipped = build_review(connection_with([]), table(), records(excel_row(), excel_row(qty=2)))
        self.assertEqual(len(review[0]['uploaded']), 2)
        self.assertEqual(review[0]['kind'], 'conflict')

    def test_keep_database_row_deletes_only_other_ids(self):
        result = self.apply('db:20', True)
        self.assertEqual(result['deleted'], 1)
        self.assertEqual(result['updated'], 0)
        statements = [call.args[0] for call in self.connection.execute.call_args_list]
        deletion = next(statement for statement in statements if str(statement).startswith('DELETE'))
        self.assertEqual(list(deletion.compile().params.values()), [[10]])
        self.connection.commit.assert_called_once()

    def test_keep_upload_preserves_survivor_nonimported_fields(self):
        result = self.apply('upload:0', True)
        self.assertEqual((result['updated'], result['deleted']), (1, 1))
        statements = [call.args[0] for call in self.connection.execute.call_args_list]
        update = next(statement for statement in statements if str(statement).startswith('UPDATE'))
        self.assertNotIn('Reference', update.compile().params)
        self.assertIn(10, update.compile().params.values())

    def test_leave_unchanged_does_not_modify_or_check_staleness(self):
        result = self.apply('skip', current=[])
        self.assertEqual(result['untouched'], 1)
        self.assertFalse(any(str(call.args[0]).startswith(('INSERT', 'UPDATE', 'DELETE')) for call in self.connection.execute.call_args_list))

    def test_delete_requires_explicit_authorization(self):
        with self.assertRaisesRegex(ValueError, 'checkbox'):
            self.apply('db:10')
        self.connection.rollback.assert_called_once()
        self.connection.commit.assert_not_called()

    def test_stale_or_invalid_choice_rolls_back(self):
        with self.assertRaisesRegex(ValueError, 'changed since preview'):
            self.apply('db:10', True, current=[dict(self.db_rows[0], Reference='edited')])
        with self.assertRaisesRegex(ValueError, 'Invalid review choice'):
            self.apply('db:999', True)
        self.connection.commit.assert_not_called()

    def test_insert_new_serial(self):
        self.group.update(snapshot=snapshot([]), existing=[])
        result = self.apply('upload:0', current=[])
        self.assertEqual(result['inserted'], 1)
        self.assertEqual(result['deleted'], 0)

    def test_history_keeps_purchase_and_rma_without_reimporting_rma(self):
        self.db_rows[0].update(entry_date='2026-04-17', invoice_number='Newegg Business')
        self.db_rows[1].update(entry_date='2026-07-20', invoice_number='Samsung', pod_number='Samsung RMA')
        self.group.update(snapshot=snapshot(self.db_rows), uploaded=[dict(self.db_rows[1], _row=1135)])
        result = self.apply('history')
        self.assertEqual((result['inserted'], result['updated'], result['deleted']), (0, 0, 0))
        self.assertFalse(any(str(call.args[0]).startswith(('INSERT', 'UPDATE', 'DELETE')) for call in self.connection.execute.call_args_list))

    def test_history_adds_only_new_events_without_deleting_or_updating(self):
        new_event = dict(self.uploaded[0], entry_date='2026-07-20', invoice_number='Samsung', pod_number='Samsung RMA')
        self.group['uploaded'] = [self.uploaded[0], new_event, dict(new_event)]
        result = self.apply('history')
        self.assertEqual((result['inserted'], result['updated'], result['deleted']), (1, 0, 0))
        statements = [str(call.args[0]) for call in self.connection.execute.call_args_list]
        self.assertEqual(sum(s.startswith('INSERT') for s in statements), 1)
        self.assertFalse(any(s.startswith(('UPDATE', 'DELETE')) for s in statements))

    def test_history_preview_marks_existing_and_new_events(self):
        new_event = dict(self.uploaded[0], entry_date='2026-07-20', invoice_number='Samsung')
        review, skipped = build_review(connection_with(self.db_rows), table(), [self.uploaded[0], new_event])
        self.assertEqual([r['_history_new'] for r in review[0]['uploaded']], [False, True])

    def test_history_rejects_stale_preview(self):
        with self.assertRaisesRegex(ValueError, 'changed since preview'):
            self.apply('history', current=[])
        self.connection.commit.assert_not_called()

    def test_write_failure_rolls_back(self):
        self.connection.execute.side_effect = [MagicMock(), MagicMock(), SQLAlchemyError('failure')]
        with self.assertRaises(SQLAlchemyError):
            self.apply('upload:0', True)
        self.connection.rollback.assert_called_once()
        self.connection.commit.assert_not_called()


if __name__ == '__main__':
    unittest.main()
