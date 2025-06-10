import re
from docx import Document
from models.models import ReceivingLog


def extract_product_details_from_word(file_path):
    try:
        document = Document(file_path)
        if not document.tables:
            return []

        table = document.tables[0]
        product_details = []
        for row in table.rows[1:-1]:  # Skip header and last row
            cells = row.cells
            if len(cells) < 4:
                continue

            sn_text = cells[2].text.strip()
            if sn_text.upper() in {"NA", "N/A", "NONE"} or not sn_text.strip():
                continue

            product_details.append({
                "product_number": cells[0].text.strip(),
                "qty": cells[1].text.strip(),
                "sn": cells[2].text.strip(),
                "notes": cells[3].text.strip()
            })
        return product_details
    except Exception as e:
        print(f"Error reading Word file {file_path}: {e}")
        return []


def get_part_name_from_db(serial_number):
    try:
        entry = ReceivingLog.query.filter_by(serial_number=serial_number).first()
        return entry.part_number if entry else None
    except Exception as e:
        print(f"DB error for {serial_number}: {e}")
        return None


def validate_sn_part_matches_via_api(file_path):
    results = []
    product_details = extract_product_details_from_word(file_path)

    for item in product_details:
        word_part = item["product_number"]
        sn_block = item["sn"]
        qty_text = item["qty"].strip()
        serials = [s.strip() for s in sn_block.split('\n') if s.strip() and s.strip().upper() not in {"NA", "N/A"}]
        sn_count = len(serials)

        try:
            expected_qty = int(qty_text)
        except ValueError:
            expected_qty = None

        qty_match = (sn_count == expected_qty) if expected_qty is not None else False

        if not serials:
            results.append({
                "serial_number": "N/A",
                "word_part": word_part,
                "db_part": None,
                "status": "❓ NOT FOUND",
                "qty_check": "N/A"
            })
            continue

        for sn in serials:
            db_part = get_part_name_from_db(sn)
            match = "✅ MATCH" if db_part == word_part else ("❌ MISMATCH" if db_part else "❓ NOT FOUND")

            results.append({
                "serial_number": sn,
                "word_part": word_part,
                "db_part": db_part,
                "status": match,
                "qty_check": f"✅ Qty OK" if qty_match else f"❌ Qty Mismatch (Expected {expected_qty}, Found {sn_count})"
            })

    return results
