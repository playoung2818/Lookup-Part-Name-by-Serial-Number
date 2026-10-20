"""Validate uploads and append only new serials to the existing receiving_log."""
import math
import hashlib
import json
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from sqlalchemy import MetaData, Table, select, text

COLUMNS = {
    'Date': 'entry_date', 'Inv#': 'invoice_number', 'Box #': 'box_number',
    'POD#': 'pod_number', 'Part#': 'part_number', 'SN#': 'serial_number', 'QTY': 'quantity',
}
STRINGS = ('invoice_number', 'box_number', 'pod_number', 'part_number', 'serial_number')
KEYS = tuple(COLUMNS.values())


def clean(value):
    return '' if value is None else str(value).strip()


def serial_key(value):
    return clean(value).upper()


def read_upload(content):
    # Bound decompression and row counts as well as the HTTP upload size.
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Workbook is too large after decompression.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('Upload a valid .xlsx or .xlsm workbook.') from exc
    try:
        if not workbook.worksheets:
            raise ValueError('Workbook must contain at least one worksheet.')
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(v).casefold() for v in next(rows, ())]
        for header in COLUMNS:
            if headers.count(header.casefold()) != 1:
                raise ValueError(f"Expected exactly one '{header}' column in the first row.")
        positions = {key: headers.index(header.casefold()) for header, key in COLUMNS.items()}
        records, errors, skipped = [], [], 0
        for number, values in enumerate(rows, 2):
            if number > 20001:
                raise ValueError('Upload at most 20,000 worksheet rows at a time.')
            row = {key: values[index] if index < len(values) else None for key, index in positions.items()}
            serial = serial_key(row['serial_number'])
            if serial in ('', 'NA', 'NAN'):
                skipped += 1
                continue
            problems = []
            for key in STRINGS:
                row[key] = clean(row[key])
                if len(row[key]) > 255:
                    problems.append(f'{key} exceeds 255 characters')
                if row[key].startswith('='):
                    problems.append(f'{key} must not be a formula')
            raw_date = row['entry_date']
            if isinstance(raw_date, datetime):
                raw_date = raw_date.date()
            elif not isinstance(raw_date, date):
                try:
                    raw_date = date.fromisoformat(clean(raw_date))
                except ValueError:
                    problems.append('Date must be an Excel date or YYYY-MM-DD')
                    raw_date = None
            row['entry_date'] = raw_date.isoformat() if raw_date else None
            try:
                quantity = 1.0 if clean(row['quantity']) == '' else float(row['quantity'])
                if not math.isfinite(quantity) or quantity <= 0:
                    raise ValueError()
                row['quantity'] = quantity
            except (ValueError, TypeError, OverflowError):
                problems.append('QTY must be a finite positive number (blank defaults to 1)')
            if not row['part_number']:
                problems.append('Part# is required')
            if problems:
                errors.append(f"Excel row {number}: " + '; '.join(problems))
            else:
                row['_row'] = number
                records.append(row)
        return records, errors, skipped
    finally:
        workbook.close()


def signature(row):
    day = row['entry_date']
    if isinstance(day, datetime):
        day = day.date()
    if isinstance(day, date):
        day = day.isoformat()
    # Serials compare case-insensitively; other identifiers retain their meaning.
    return (day, *(serial_key(row[key]) if key == 'serial_number' else clean(row[key]) for key in STRINGS), row['quantity'])


def receiving_table(connection):
    # Reflection intentionally requires an existing table; never creates schema.
    table = Table('receiving_log', MetaData(), autoload_with=connection)
    if not set(KEYS).issubset(table.c.keys()):
        raise ValueError('receiving_log is missing required columns.')
    return table


def classify(connection, table, records):
    existing = {}
    for row in connection.execute(select(*(table.c[key] for key in KEYS))).mappings():
        key = serial_key(row['serial_number'])
        if key not in ('', 'NA', 'NAN'):
            existing.setdefault(key, []).append(row)
    pending, seen, errors, skipped = [], {}, [], 0
    for row in records:
        key = serial_key(row['serial_number'])
        if key in seen:
            if signature(seen[key]) == signature(row):
                skipped += 1
            else:
                errors.append(f"Excel row {row['_row']}: serial {row['serial_number']} conflicts with another uploaded row.")
            continue
        seen[key] = row
        matches = existing.get(key, [])
        if len(matches) > 1:
            errors.append(f"Serial {row['serial_number']} already has {len(matches)} database records; resolve duplicates first.")
        elif matches:
            if signature(matches[0]) == signature(row):
                skipped += 1
            else:
                errors.append(f"Excel row {row['_row']}: serial {row['serial_number']} already exists with different receiving details.")
        else:
            pending.append(row)
    return pending, errors, skipped


def lock_receiving(connection):
    if connection.dialect.name == 'sqlite':
        connection.exec_driver_sql('BEGIN IMMEDIATE')
    elif connection.dialect.name == 'postgresql':
        connection.execute(text("SET LOCAL lock_timeout = '10s'"))
        connection.execute(text('LOCK TABLE receiving_log IN SHARE ROW EXCLUSIVE MODE'))
    else:
        raise ValueError('Safe receiving imports require PostgreSQL or SQLite.')


def import_records(engine, records):
    with engine.connect() as connection:
        try:
            lock_receiving(connection)
            table = receiving_table(connection)
            pending, errors, skipped = classify(connection, table, records)
            if errors:
                raise ValueError('Import cancelled: ' + ' '.join(errors[:20]))
            if pending:
                data = []
                for row in pending:
                    item = {key: row[key] for key in KEYS}
                    item['entry_date'] = date.fromisoformat(item['entry_date'])
                    data.append(item)
                connection.execute(table.insert(), data)
            connection.commit()
            return len(pending), skipped
        except Exception:
            connection.rollback()
            raise


def snapshot(rows):
    """Hash every database field, including fields not supplied by Excel."""
    packed = [json.dumps(dict(row), sort_keys=True, default=str) for row in rows]
    return hashlib.sha256(json.dumps(sorted(packed)).encode()).hexdigest()


def database_groups(connection, table):
    if 'id' not in table.c:
        raise ValueError('Review requires an existing receiving_log.id column.')
    groups = {}
    ids = set()
    for row in connection.execute(select(table)).mappings():
        row = dict(row)
        if row['id'] is None or row['id'] in ids:
            raise ValueError('receiving_log.id must contain unique, non-null IDs for safe review.')
        ids.add(row['id'])
        key = serial_key(row['serial_number'])
        if key not in ('', 'NA', 'NAN'):
            groups.setdefault(key, []).append(row)
    return groups


def build_review(connection, table, records):
    existing = database_groups(connection, table)
    uploaded = {}
    skipped = 0
    for row in records:
        candidates = uploaded.setdefault(serial_key(row['serial_number']), [])
        if any(signature(candidate) == signature(row) for candidate in candidates):
            skipped += 1
        else:
            candidates.append(row)
    review = []
    for key, candidates in uploaded.items():
        matches = sorted(existing.get(key, []), key=lambda row: row['id'])
        if len(matches) == 1 and len(candidates) == 1 and signature(matches[0]) == signature(candidates[0]):
            skipped += 1
            continue
        # History mode retains every existing event and only adds distinct ones.
        known = {signature(row) for row in matches}
        candidates = [dict(row, _history_new=signature(row) not in known) for row in candidates]
        # Store only affected serials, not the thousands of unchanged workbook rows.
        review.append({'key': key, 'uploaded': candidates, 'existing': matches,
                       'snapshot': snapshot(matches),
                       'kind': 'new' if not matches and len(candidates) == 1 else 'conflict'})
    return review, skipped


def apply_review(engine, review, choices, allow_delete=False):
    """Apply explicit per-serial decisions atomically; stale previews never write."""
    counts = {'inserted': 0, 'updated': 0, 'deleted': 0, 'untouched': 0}
    with engine.connect() as connection:
        try:
            lock_receiving(connection)
            table = receiving_table(connection)
            current = database_groups(connection, table)
            plans = []
            for index, group in enumerate(review):
                choice = choices.get(str(index), 'skip')
                if choice == 'skip':
                    counts['untouched'] += 1
                    continue
                matches = sorted(current.get(group['key'], []), key=lambda row: row['id'])
                if snapshot(matches) != group['snapshot']:
                    raise ValueError(f"Serial {group['key']} changed since preview. Nothing saved; upload again.")
                if choice == 'history':
                    known = {signature(row) for row in matches}
                    for row in group['uploaded']:
                        event = signature(row)
                        if event not in known:
                            plans.append(('upload:history', row, None, []))
                            known.add(event)
                    continue
                options = {f"db:{row['id']}": row for row in matches}
                options.update({f'upload:{i}': row for i, row in enumerate(group['uploaded'])})
                if choice not in options:
                    raise ValueError('Invalid review choice. Nothing saved.')
                chosen = options[choice]
                # Upload replaces the lowest-ID existing row, preserving Reference
                # and all other non-imported fields on that survivor.
                survivor = chosen if choice.startswith('db:') else (matches[0] if matches else None)
                removed = [row['id'] for row in matches if survivor is None or row['id'] != survivor['id']]
                if removed and not allow_delete:
                    raise ValueError('Confirm the database-row deletion checkbox before resolving duplicates.')
                plans.append((choice, chosen, survivor, removed))
            for choice, chosen, survivor, removed in plans:
                if removed:
                    result = connection.execute(table.delete().where(table.c.id.in_(removed)))
                    if result.rowcount != len(removed):
                        raise ValueError('Database rows changed unexpectedly. Nothing saved.')
                    counts['deleted'] += len(removed)
                if choice.startswith('upload:'):
                    data = {key: chosen[key] for key in KEYS}
                    data['entry_date'] = date.fromisoformat(data['entry_date'])
                    if survivor is None:
                        connection.execute(table.insert().values(**data))
                        counts['inserted'] += 1
                    else:
                        result = connection.execute(table.update().where(table.c.id == survivor['id']).values(**data))
                        if result.rowcount != 1:
                            raise ValueError('Database survivor changed unexpectedly. Nothing saved.')
                        counts['updated'] += 1
            connection.commit()
            return counts
        except Exception:
            connection.rollback()
            raise
